"""Small python-docx helpers for the research report (US Letter, compact academic style)."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import pandas as pd

INK = RGBColor(0x14, 0x17, 0x1C); MUTED = RGBColor(0x4A, 0x52, 0x60); ACCENT = RGBColor(0x1F, 0x4E, 0x8C)


def new_doc():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.left_margin = s.right_margin = Inches(0.85); s.top_margin = s.bottom_margin = Inches(0.85)
    st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.paragraph_format.space_after = Pt(4); st.paragraph_format.line_spacing = 1.08
    for lvl, size in [(1, 13.5), (2, 11.5), (3, 10.5)]:
        h = d.styles[f"Heading {lvl}"]; h.font.name = "Calibri"; h.font.size = Pt(size); h.font.bold = True
        h.font.color.rgb = ACCENT if lvl == 1 else INK
        h.paragraph_format.space_before = Pt(10 if lvl == 1 else 6); h.paragraph_format.space_after = Pt(3)
    # page numbers in footer
    p = s.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field(p, "PAGE")
    return d


def _field(paragraph, code):
    r = paragraph.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = code
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    r._r.append(f1); r._r.append(it); r._r.append(f2)
    r.font.size = Pt(9); r.font.color.rgb = MUTED


def para(d, text, size=None, italic=False, color=None, align=None, space_after=None, bold_lead=None):
    """Paragraph with **bold** inline markup support."""
    p = d.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead + " "); r.bold = True
        if size: r.font.size = Pt(size)
    _runs(p, text, size=size, italic=italic, color=color)
    if align == "center": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if align == "justify": p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if space_after is not None: p.paragraph_format.space_after = Pt(space_after)
    return p


def _runs(p, text, size=None, italic=False, color=None):
    """**bold** and *italic* inline markup."""
    import re
    for i, part in enumerate(text.split("**")):
        if not part: continue
        for j, sub in enumerate(re.split(r"(?<![\w*])\*(?!\s)([^*]+?)(?<!\s)\*(?![\w*])", part)):
            if not sub: continue
            r = p.add_run(sub); r.bold = (i % 2 == 1); r.italic = italic or (j % 2 == 1)
            if size: r.font.size = Pt(size)
            if color: r.font.color.rgb = color


def bullets(d, items, size=None):
    for it in items:
        p = d.add_paragraph(style="List Bullet")
        _runs(p, it, size=size)
        p.paragraph_format.space_after = Pt(2)


def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcolor); tcPr.append(sh)


def table(d, df, caption=None, widths=None, size=8.5, fmt=None, first_col_left=True, note=None):
    if caption:
        c = para(d, caption, size=9.5, space_after=2); c.runs[0].bold = True if c.runs else None
    t = d.add_table(rows=1, cols=len(df.columns)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, col in enumerate(df.columns):
        cell = t.rows[0].cells[j]; cell.text = ""; r = cell.paragraphs[0].add_run(str(col)); r.bold = True; r.font.size = Pt(size)
        _shade(cell, "E8EDF4")
    for _, row in df.iterrows():
        cells = t.add_row().cells
        for j, col in enumerate(df.columns):
            v = row[col]
            if fmt and col in fmt and isinstance(v, (int, float)) and pd.notna(v): s = fmt[col].format(v)
            elif isinstance(v, float): s = "–" if pd.isna(v) else f"{v:,.2f}"
            else: s = str(v)
            cells[j].text = ""; r = cells[j].paragraphs[0].add_run(s); r.font.size = Pt(size)
            cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if (j == 0 and first_col_left) else WD_ALIGN_PARAGRAPH.RIGHT
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs: p.paragraph_format.space_after = Pt(0)
    total = 6.7; n = len(df.columns)
    if not widths:
        first = min(2.6, max(1.4, total * 0.35)) if n > 2 else total / n
        widths = [first] + [(total - first) / (n - 1)] * (n - 1) if n > 1 else [total]
    t.autofit = False
    for j, w in enumerate(widths):
        t.columns[j].width = Inches(w)
        for row in t.rows: row.cells[j].width = Inches(w)
    if note: para(d, note, size=8, italic=True, color=MUTED, space_after=6)
    else: d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(d, path, caption, width=6.6):
    from pathlib import Path
    if not Path(path).exists():
        para(d, f"[Figure pending: {Path(path).name}]", size=9, italic=True, color=MUTED); return
    d.add_picture(str(path), width=Inches(width)); d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(d, caption, size=8.5, italic=True, color=MUTED, space_after=6)
